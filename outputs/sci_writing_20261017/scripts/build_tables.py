"""Prepare manuscript tables from frozen outputs; no statistical model is run.

Run with the existing bundled Python environment. XLSX authoring is performed
separately by tables/build_workbook.mjs using the bundled artifact-tool API.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import json
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "outputs/sci_writing_20261017"
PREV = ROOT / "outputs/locus_novelty_20261017"
TABLES = OUT / "tables"
MAPS = OUT / "source_data"
REVIEWS = OUT / "reviews"
for directory in (TABLES, MAPS, REVIEWS):
    directory.mkdir(parents=True, exist_ok=True)


def read_tsv(path: Path) -> list[dict]:
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def write_tsv(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def save_json(path: Path, payload) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def relative(path: Path) -> str:
    return str(path.relative_to(ROOT))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


INPUT_PATHS = [
    ROOT / "config/traits.tsv",
    ROOT / "config/analysis.yml",
    ROOT / "codes/20_run_t06_gwas_gwas_susie_coloc.R",
    ROOT / "documentation/analysis/PHENOTYPE_AND_CONTROL_PROTOCOL.md",
    ROOT / "outcome_data/t06/gwas_gwas_susie_coloc.tsv",
    ROOT / "outcome_data/t06/susie_signal_catalog.tsv",
    ROOT / "outcome_data/t06/susie_credible_set_members.tsv.gz",
    ROOT / "outcome_data/t07/genetic_signal_five_classification.tsv",
    ROOT / "outcome_data/t11b/p1a_feature_results.tsv",
    ROOT / "outcome_data/t16/sources/GCST90018635.json",
    ROOT / "outcome_data/t16/sources/GCST90018635_meta.yaml",
    ROOT / "outcome_data/t17/sources/hum0014_sample_size.xlsx",
    ROOT / "outcome_data/t17/sources/hum0014_README.txt",
    ROOT / "documentation/analysis/T18_BBJ_INPUT_APPLICABILITY_AUDIT.md",
    ROOT / "documentation/analysis/T18_BBJ_ORIGINAL_AIHT_FULL_QC.md",
    PREV / "manuscript_blueprint.md",
    PREV / "locus_increment_matrix.tsv",
    PREV / "review_reeve_bujnis/evidence_matches.tsv",
    PREV / "review_reeve_bujnis/sources/reeve_tables.xlsx",
    PREV / "review_reeve_bujnis/sources/bujnis_tables.xlsx",
    PREV / "review_multiome/published_target_disease_hits.tsv",
    PREV / "review_multiome/supp_table_26.xlsx",
    PREV / "review_multiome/source_manifest.json",
    PREV / "root_sources/source_notes.md",
    PREV / "root_sources/makkonen.html",
]
input_hashes = {relative(path): sha256(path) for path in INPUT_PATHS}
checks = []


def check(name: str, passed: bool, detail) -> None:
    checks.append({"name": name, "passed": bool(passed), "detail": detail})


traits = {row["trait_id"]: row for row in read_tsv(ROOT / "config/traits.tsv")}
G = "E4_GRAVES_STRICT"
A = "E4_HYTHY_AI_STRICT"
pairs = read_tsv(ROOT / "outcome_data/t06/gwas_gwas_susie_coloc.tsv")
all_units = read_tsv(ROOT / "outcome_data/t07/genetic_signal_five_classification.tsv")
units = {row["signal_unit_id"]: row for row in all_units}
inventory = read_tsv(PREV / "locus_increment_matrix.tsv")
features = read_tsv(ROOT / "outcome_data/t11b/p1a_feature_results.tsv")
source_matches = read_tsv(PREV / "review_reeve_bujnis/evidence_matches.tsv")
nature_hits = read_tsv(PREV / "review_multiome/published_target_disease_hits.tsv")
members = read_tsv(ROOT / "outcome_data/t06/susie_credible_set_members.tsv.gz")
catalog = read_tsv(ROOT / "outcome_data/t06/susie_signal_catalog.tsv")

check("FinnGen_primary_counts", [(int(traits[t]["n_cases"]), int(traits[t]["n_controls"])) for t in (G, A)] == [(4247, 495939), (56098, 358456)], "config/traits.tsv primary rows")
check("pair_and_region_denominators", len(pairs) == 42 and len({r["pair_id"] for r in pairs}) == 16 and len({r["graves_region"] for r in pairs}) == 15 and len({r["aiht_region"] for r in pairs}) == 16, "42 component comparisons across 16 region pairs, 15 unique Graves regions and 16 unique AIHT regions")
check("T07_units_and_prioritized_rows", len(all_units) == 304 and len(inventory) == 10 and len({r["signal_unit_id"] for r in inventory}) == 10, "Ten prioritized discovery candidates are a selected subset, not ten independent loci or a prespecified genome-wide denominator")
check("QTL_selection_denominator", len(features) == 26 and len({r["locus_id"] for r in features}) == 7, "26 selected feature rows across 7 loci")

bbj_catalog = json.loads((ROOT / "outcome_data/t16/sources/GCST90018635.json").read_text())
check("BBJ_Hashimoto_source_identity", bbj_catalog["accessionId"] == "GCST90018635" and bbj_catalog["cohort"] == "BBJ" and bbj_catalog["initialSampleSize"] == "537 East Asian ancestry cases, 172,656 East Asian ancestry controls", bbj_catalog["initialSampleSize"])
bbj_sample_wb = openpyxl.load_workbook(ROOT / "outcome_data/t17/sources/hum0014_sample_size.xlsx", read_only=True, data_only=True)
bbj_gd_row = list(bbj_sample_wb["Sample size"].iter_rows(min_row=24, max_row=24, values_only=True))[0]
check("BBJ_Graves_sample_workbook", bbj_gd_row[2] == "hum0014.v17.GD.v1" and bbj_gd_row[3] == 2176 and bbj_gd_row[6] == 210277 and bbj_gd_row[3] + bbj_gd_row[6] == 212453, {"sheet": "Sample size", "row": 24, "values": list(bbj_gd_row)})
bbj_sample_wb.close()
check("BBJ_Graves_original_build", "Positions are based on hg19" in (ROOT / "outcome_data/t17/sources/hum0014_README.txt").read_text(), "Original README, physical line 14")
check("BBJ_Hashimoto_original_build", "genome_assembly: GRCh37" in (ROOT / "outcome_data/t16/sources/GCST90018635_meta.yaml").read_text(), "Original released metadata, physical line 17")

table1_columns = [
    {"key": "data", "label": "Data or comparison"},
    {"key": "definition", "label": "Phenotype or feature definition"},
    {"key": "sample_unit", "label": "Sample size or comparison unit"},
    {"key": "role_limit", "label": "Role and comparability"},
]

table1_rows = [
    {"row_id": "FG_GD", "section": "Primary discovery", "data": "FinnGen R13 Graves’ disease\nE4_GRAVES_STRICT", "definition": "(Diffuse-goitre thyrotoxicosis OR Graves’ ophthalmopathy) AND treatment-support condition AND no amiodarone. The support condition uses carbimazole purchase or total thyroidectomy.", "sample_unit": f"{int(traits[G]['n_cases']):,} cases; {int(traits[G]['n_controls']):,} controls", "role_limit": "Finnish-enriched European ancestry; GRCh38. Controls are noncases without additional thyroid or autoimmune exclusions."},
    {"row_id": "FG_AIHT", "section": "Primary discovery", "data": "FinnGen R13 autoimmune hypothyroidism\nE4_HYTHY_AI_STRICT", "definition": "E4_HYTHYNAS AND NOT E4_HYTHY_AI_EXCL. The source endpoint combines diagnostic, reimbursement (KELA 104) and thyroid-hormone records under its event rules.", "sample_unit": f"{int(traits[A]['n_cases']):,} cases; {int(traits[A]['n_controls']):,} controls", "role_limit": "Finnish-enriched European ancestry; GRCh38. Controls exclude AUTOIMMUNE and E4_HYTHY_AI_EXCL. Prior Graves’ disease and radioiodine are not explicitly excluded by the public rule."},
    {"row_id": "COMPARABILITY", "section": "Primary discovery", "data": "Cross-endpoint comparability", "definition": "Summary statistics only. The endpoints do not establish mutually exclusive lifetime diagnoses or a common control definition.", "sample_unit": "Exact individual overlap unavailable", "role_limit": "Cross-trait association components are compared. Disease transition and person-level case–case effects cannot be recovered."},
    {"row_id": "CS_COMPARISON", "section": "Comparison units", "data": "Official FinnGen R13 SuSiE components", "definition": "Pairwise comparisons of detected components within overlapping fine-mapping regions, with quality and common-variant coverage retained.", "sample_unit": "42 component comparisons; 16 region pairs (15 unique Graves regions; 16 AIHT regions)", "role_limit": "Components and region pairs are analysis units. An H3 comparison does not establish absence of a Graves counterpart for an AIHT component."},
    {"row_id": "CANDIDATES", "section": "Comparison units", "data": "Prioritized discovery variants", "definition": "Ten candidates selected from the existing discovery classification; marginal associations aligned to the displayed ALT allele.", "sample_unit": "10 displayed candidates from 304 constructed analysis units", "role_limit": "The display is selected after discovery. The 304 constructed units are not 304 independent biological mechanisms."},
    {"row_id": "QTL", "section": "Supporting evidence", "data": "FinnGen immune multiome v1", "definition": "Chromatin-accessibility and expression QTLs from healthy-donor peripheral blood mononuclear cells; selected cell–feature comparisons.", "sample_unit": "26 selected features across 7 loci", "role_limit": "Regulatory context. The same QTL resource and overlapping R12/R13 FinnGen GWAS lineage preclude an independent replication claim [Kanai 2026]."},
    {"row_id": "BBJ_GD", "section": "Supporting evidence", "data": "BioBank Japan Graves’ disease\nhum0014.v17.GD.v1 (2019)", "definition": "Japanese Graves’ disease GWAS, both-sex autosomal release; used for the frozen representative-variant checks.", "sample_unit": "2,176 cases; 210,277 controls", "role_limit": "Japanese ancestry; GRCh37 source coordinates. External variant-level evidence only; release-matched signed LD and cross-disease covariance are unavailable."},
    {"row_id": "BBJ_HD", "section": "Supporting evidence", "data": "BioBank Japan Hashimoto thyroiditis\nGCST90018635 / hum0197.v3.BBJ.HD.v1", "definition": "BBJ-only Hashimoto thyroiditis, mapped to AIHT for the frozen representative-variant checks [Sakaue 2021].", "sample_unit": "537 cases; 172,656 controls", "role_limit": "East Asian ancestry, Japan; GRCh37 source coordinates. The Catalog file and original BBJ archive represent the same study, not two replications. BBJ cross-trait sample overlap is unknown."},
]

table1_notes = [
    "AIHT, autoimmune hypothyroidism; BBJ, BioBank Japan; caQTL, chromatin-accessibility quantitative trait locus; CS, credible set; eQTL, expression quantitative trait locus; GWAS, genome-wide association study; LD, linkage disequilibrium; QTL, quantitative trait locus; SuSiE, sum of single effects.",
    "FinnGen case and control counts refer to the released GWAS samples, not the larger pre-genotype-quality-control endpoint counts. The summary endpoint labelled AIHT does not verify autoimmune aetiology in each individual and is not an antibody- or pathology-confirmed Hashimoto cohort.",
    "For E4_HYTHY_AI_EXCL, the public diagnosis expression includes E05.1–E05.9, but not E05.0 (Graves’ disease); the surgical expression explicitly covers total thyroidectomy (BAA60), and no radioiodine exclusion is shown. The study could not apply additional person-level exclusions.",
    "FinnGen GWAS and fine-mapping results are reused outputs. The immune multiome resource provides regulatory context. The 26 features were selected before the existing targeted QTL checks, and their findings cannot estimate a genome-wide confirmation rate.",
    "BBJ phenotype mappings are approximate. The published Graves and Hashimoto outcomes do not reproduce the FinnGen strict endpoint definitions. Existing BBJ tests concern representative marginal associations; they do not revalidate component identity, signal classes or cross-disease causal effects.",
    "The direct cross-trait colocalization model does not explicitly account for correlated estimation errors between the two FinnGen endpoint GWAS.",
]

# Each entry is a source-bounded comparison, not an inference of new association.
prior_text = {
    "GSU018": "Concordant associations at rs2476601 (Reeve 2026); published index rs6679677 is in current Graves CS1 (Bujnis 2026).",
    "GSU042": "Concordant associations at rs182482858, annotated to ZAP70 (Reeve 2026).",
    "GSU052": "rs11571297 associations and QTL evidence (Reeve 2026; Kanai 2026); regional complexity and multiple CSs (Wallace 2013; Bujnis 2026).",
    "GSU056": "Opposite marginal associations at rs6724363 (Reeve 2026); the AIHT component is also in Kanai 2026.",
    "GSU073": "Published index rs11335233 belongs to both current CS1s (Reeve 2026); five regional CSs were reported by Bujnis 2026.",
    "GSU124": "Both disease components at rs748670681, regulatory annotations and the same R12 cell–feature evidence (Kanai 2026).",
    "GSU138": "Graves rs2466074 (Kanai 2026); published index rs113350646 is in current AIHT CS1 and has opposite marginal associations (Reeve 2026).",
    "GSU174": "AIHT rs35508742 (Kanai 2026); regional associations and multiple components (Reeve 2026; Bujnis 2026). Reeve rs71508903 is the current AIHT CS1 lead.",
    "GSU231": "Graves rs1023586 and previously separated TSHR associations (Makkonen 2024; Kanai 2026); published index rs35277102 belongs to current AIHT CS1 (Reeve 2026).",
    "GSU254": "Published index rs5818256 is in current AIHT CS1 and has opposite associations (Reeve 2026); prior CS2 members overlap both current CSs (Bujnis 2026).",
}
match_text = {
    "GSU018": "Exact variant; additional current-CS membership",
    "GSU042": "Exact variant",
    "GSU052": "Exact variant; regional/multiple-CS evidence",
    "GSU056": "Exact variant; opposite marginal associations already reported",
    "GSU073": "Published index in both current CSs; different representative",
    "GSU124": "Exact variant; source regulatory evidence inherited",
    "GSU138": "Exact Graves variant; different index in AIHT CS1",
    "GSU174": "Exact AIHT variant; regional/multiple-CS evidence",
    "GSU231": "Exact Graves variant; different index in AIHT CS1",
    "GSU254": "Published index in AIHT CS1; other members in both CSs",
}
increment_text = {
    "GSU018": "Maps a known concordant association to the detected cross-trait pair and distinguishes other AIHT components.",
    "GSU042": "Adds the detected component correspondence. TMEM131/ZAP70 annotation differences do not identify a new locus or causal gene.",
    "GSU052": "Resolves the tested Graves CS1 correspondence to AIHT CS1 while retaining other detected components and five unresolved comparisons. Cross-study CS labels cannot be equated by number.",
    "GSU056": "Adds a shared-model comparison to a previously reported opposite marginal association. Marginal direction is not a conditional causal effect.",
    "GSU073": "Maps existing regional components. A different representative SNP does not establish a new independent signal.",
    "GSU124": "Adds a cross-trait component pairing while declaring inherited regulatory evidence. It does not establish a new independently replicated mechanism.",
    "GSU138": "Shows a distinct-component comparison and sensitivity of the marginal Graves-predominant label to the equivalence bound; it does not imply gene-wide specificity.",
    "GSU174": "Maps detected Graves CS1 to AIHT CS2, rather than to AIHT CS1/CS3. The H4 representative is not either full-region lead; cross-study CS equivalence remains incomplete.",
    "GSU231": "Updates the R13 pairing and reports stable representative-variant classification across the tested bounds. Separate TSHR associations were already reported.",
    "GSU254": "Adds the current shared-model pairing to prior regional opposite marginal associations; the opposite-direction mechanism is not established.",
}
label_text = {
    "GSU018": "PTPN22", "GSU042": "TMEM131/ZAP70", "GSU052": "CTLA4/ICOS",
    "GSU056": "AC007563.2/DIRC3", "GSU073": "LPP", "GSU124": "TNRC18",
    "GSU138": "NRG1", "GSU174": "ARID5B/RTKN2", "GSU231": "TSHR", "GSU254": "LINC01229/MAFTRR",
}


def comparison_text(rows: list[dict]) -> str:
    result = []
    for label, cls in (("H4", "shared_strong_H4"), ("H3", "distinct_signals_H3")):
        keys = [f"G{r['graves_cs']}–A{r['aiht_cs']}" for r in rows if r["classification"] == cls]
        if keys:
            result.append(f"{label}: " + ", ".join(keys))
    unresolved = [f"G{r['graves_cs']}–A{r['aiht_cs']}" for r in rows if r["classification"].startswith("unresolved")]
    if unresolved:
        result.append("Unresolved: " + ", ".join(unresolved))
    return ". ".join(result) + "."


table2_columns = [
    {"key": "candidate", "label": "Candidate, region pair and representative variant"},
    {"key": "prior_evidence", "label": "Prior evidence"},
    {"key": "match_level", "label": "Evidence match level"},
    {"key": "comparison", "label": "Current component comparisons"},
    {"key": "increment", "label": "Information added and interpretation"},
]
table2_rows = []
for row in inventory:
    gsu = row["signal_unit_id"]
    unit = units[gsu]
    check("candidate_identity_" + gsu, unit["representative_variant"] == row["representative_GRCh38"] and unit["representative_reported_rsids"] == row["rsid"], {"variant": unit["representative_variant"], "rsid": unit["representative_reported_rsids"]})
    table2_rows.append({"row_id": gsu, "candidate": f"{label_text[gsu]}\n{gsu}; {row['pair_id']}\n{row['rsid']}\n{row['representative_GRCh38']}", "prior_evidence": prior_text[gsu], "match_level": match_text[gsu], "comparison": comparison_text([p for p in pairs if p["pair_id"] == row["pair_id"]]), "increment": increment_text[gsu]})

table2_notes = [
    "The table contains ten prioritized discovery candidates; it is not a prespecified set selected before discovery. Region names are positional labels and do not establish causal genes. All listed regions have prior association evidence; no new independent risk locus is established by this comparison.",
    "G and A identify Graves and AIHT components, respectively, within the indicated current FinnGen R13 region pair (for example, G1–A2 means Graves CS1 compared with AIHT CS2). The same Graves CS1 appears in every supported H4/H3 comparison in the six mixed regions (RP01, RP02, RP06, RP07, RP10 and RP12). Counts of such comparisons are not counts of independent Graves signals.",
    "H4 denotes the existing quality-qualified shared-variant model classification; H3 denotes the existing quality-qualified distinct-variant model classification. Unresolved comparisons retain the existing low-purity or insufficient-common-variant-coverage status. An H3 result does not establish disease specificity, a zero effect in the other disease or absence of an undetected counterpart.",
    "Exact variant denotes a match in genome build, position, reference allele and alternate allele; rsIDs are retained as identifiers. An rsID-only match is identified separately and does not by itself establish exact allele identity. CS membership means that a different published index variant is in the current 95% CS; it does not prove cross-study causal-signal identity. Region-level evidence is broader still. CS numbers in other studies belong to those studies’ models and are not interchangeable with the current labels.",
    "ST denotes an original supplementary table; row numbers are physical Excel worksheet rows, including header rows. At LPP, the five prior CSs are in Bujnis ST4; the detailed row-level evidence map accompanies this table. Kanai 2026 uses the FinnGen immune multiome resource and R12 disease data, so the cited regulatory evidence is inherited source context rather than independent confirmation of the present R13 reanalysis.",
    "At ARID5B/RTKN2, the current full-region leads are chr10_62212113_A_T (Graves CS1) and chr10_62286595_T_G (AIHT CS2); rs35508742 is the highest-posterior H4 representative for their comparison. Marginal allelic directions and equivalence-bound labels concern representative variants and do not give signal-specific conditional causal directions.",
]

references = [
    {"id": "Reeve 2026", "citation": "Reeve et al. (2026). Genome-wide association analyses of autoimmune hypothyroidism reveal autoimmune and thyroid-specific contributions and an inverse relationship with cancer risk. Nature Genetics.", "doi": "10.1038/s41588-026-02521-1", "url": "https://www.nature.com/articles/s41588-026-02521-1"},
    {"id": "Bujnis 2026", "citation": "Bujnis et al. (2026). Hashimoto thyroiditis GWAS and fine-mapping study. Nature Genetics. The complete publisher title is retained in the accompanying source-reference metadata.", "doi": "10.1038/s41588-026-02704-w", "url": "https://www.nature.com/articles/s41588-026-02704-w"},
    {"id": "Kanai 2026", "citation": "Kanai et al. (2026). Population-scale immune multiome atlas reveals regulatory disease mechanisms. Nature.", "doi": "10.1038/s41586-026-11078-2", "url": "https://www.nature.com/articles/s41586-026-11078-2"},
    {"id": "Wallace 2013", "citation": "Wallace (2013). Statistical testing of shared genetic control for potentially related traits. Genetic Epidemiology.", "doi": "10.1002/gepi.21765", "url": "https://pmc.ncbi.nlm.nih.gov/articles/PMC4158901/"},
    {"id": "Makkonen 2024", "citation": "Makkonen et al. (2024). Mechanisms of thyrotropin receptor–mediated phenotype variability deciphered by gene mutations and M453T-knockin model. JCI Insight.", "doi": "10.1172/jci.insight.167092", "url": "https://insight.jci.org/articles/view/167092"},
    {"id": "Sakaue 2021", "citation": "Sakaue et al. (2021). A cross-population atlas of genetic associations for 220 human phenotypes. Nature Genetics.", "doi": "10.1038/s41588-021-00931-x", "url": "https://www.nature.com/articles/s41588-021-00931-x"},
]

# Resolve the exact Bujnis title from the original local publisher metadata.
bujnis_html = (PREV / "review_reeve_bujnis/sources/bujnis_publisher.html").read_text()
import html
bujnis_title_match = re.search(r'<meta name="citation_title" content="([^"]+)"', bujnis_html)
if not bujnis_title_match:
    raise ValueError("Original Bujnis publisher citation title unavailable")
bujnis_title = html.unescape(bujnis_title_match.group(1))
references[1]["citation"] = f"Bujnis et al. (2026). {bujnis_title}. Nature Genetics."
check("Bujnis_exact_publisher_title", bool(bujnis_title), bujnis_title)
makkonen_html = (PREV / "root_sources/makkonen.html").read_text()
check("Makkonen_publication_year", bool(re.search(r'(citation_date|citation_publication_date)[^>]+2024', makkonen_html)), "Original JCI metadata dates the article to 2024")

source_maps = {"table1": [], "table2": []}


def add_map(table, row_id, column, source, locator, evidence_key, relation):
    source_maps[table].append({"table": table, "row_id": row_id, "column": column, "source_file": relative(source), "source_locator": locator, "evidence_key": evidence_key, "relation": relation})


protocol = ROOT / "documentation/analysis/PHENOTYPE_AND_CONTROL_PROTOCOL.md"
protocol_lines = protocol.read_text().splitlines()


def line_for(text):
    hits = [str(i + 1) for i, line in enumerate(protocol_lines) if text in line]
    if not hits:
        raise ValueError(f"Missing protocol text: {text}")
    return ",".join(hits)


for row_id, trait, config_line, predicate in (("FG_GD", G, 2, "AND GRAVES_COND"), ("FG_AIHT", A, 3, "E4_HYTHYNAS")):
    add_map("table1", row_id, "sample_unit", ROOT / "config/traits.tsv", f"physical line {config_line}", trait, "exact case/control counts")
    for column in ("data", "definition", "role_limit"):
        add_map("table1", row_id, column, protocol, "physical lines " + line_for(predicate), trait, "official endpoint audit; full associated section applies")
    add_map("table1", row_id, "role_limit", protocol, "physical lines " + line_for("GRCh38"), trait, "build and ancestry audit")
add_map("table1", "FG_GD", "role_limit", protocol, "physical lines " + line_for("对照为所有非病例"), G, "control definition")
add_map("table1", "FG_AIHT", "role_limit", protocol, "physical lines " + line_for("NOT AUTOIMMUNE"), A, "control exclusions")
for col in ("definition", "sample_unit", "role_limit"):
    add_map("table1", "COMPARABILITY", col, protocol, "physical lines " + line_for("Exact individual overlap"), G + " vs " + A, "individual-level overlap not available")
for col in ("definition", "sample_unit", "role_limit"):
    add_map("table1", "CS_COMPARISON", col, ROOT / "outcome_data/t06/gwas_gwas_susie_coloc.tsv", "all 42 data rows; group by pair_id and source-region strings", "42 comparisons / 16 pairs", "enumeration of existing records only")
    add_map("table1", "CANDIDATES", col, ROOT / "outcome_data/t07/genetic_signal_five_classification.tsv", "all 304 data rows; prioritized IDs in table2", "304 constructed units; 10 candidates", "existing classification and accepted selection")
    add_map("table1", "QTL", col, ROOT / "outcome_data/t11b/p1a_feature_results.tsv", "all 26 data rows", "26 selected features / 7 locus_id values", "existing targeted subset only")
add_map("table1", "QTL", "role_limit", PREV / "review_multiome/tnrc18_same_cell_feature_trait_R12.tsv", "all 4 data rows with original source-path and source-line locators", "GSU124 R12 caQTL/eQTL", "source lineage overlap")
for col in ("data", "definition", "sample_unit", "role_limit"):
    add_map("table1", "BBJ_GD", col, ROOT / "documentation/analysis/T18_BBJ_INPUT_APPLICABILITY_AUDIT.md", "source/cohort contract, physical lines 15–25", "hum0014.v17.GD.v1", "frozen input identity and use limits")
    add_map("table1", "BBJ_HD", col, ROOT / "outcome_data/t16/sources/GCST90018635.json", "accessionId; cohort; initialSampleSize; ancestries; diseaseTrait", "GCST90018635", "original Catalog metadata")
add_map("table1", "BBJ_GD", "sample_unit", ROOT / "outcome_data/t17/sources/hum0014_sample_size.xlsx", "Sample size!A24:I24", "Graves’ disease", "original released sample counts")
add_map("table1", "BBJ_GD", "role_limit", ROOT / "outcome_data/t17/sources/hum0014_README.txt", "physical line 14", "Positions are based on hg19", "original source genome assembly")
add_map("table1", "BBJ_HD", "role_limit", ROOT / "outcome_data/t16/sources/GCST90018635_meta.yaml", "physical line 17", "genome_assembly: GRCh37", "original source genome assembly")
add_map("table1", "BBJ_HD", "role_limit", ROOT / "documentation/analysis/T18_BBJ_ORIGINAL_AIHT_FULL_QC.md", "original archive identity and GCST reconciliation sections", "hum0197.v3.BBJ.HD.v1", "same underlying BBJ study; not independent replication")
add_map("table1", "note_6", "notes", ROOT / "codes/20_run_t06_gwas_gwas_susie_coloc.R", "physical lines 316–325", "primary and sensitivity coloc_bf_bf calls", "direct model calls do not supply cross-trait estimation-error covariance")

for row in inventory:
    gsu, pair_id = row["signal_unit_id"], row["pair_id"]
    unit_line = next(i + 2 for i, u in enumerate(all_units) if u["signal_unit_id"] == gsu)
    for col in ("candidate",):
        add_map("table2", gsu, col, ROOT / "outcome_data/t07/genetic_signal_five_classification.tsv", f"physical line {unit_line}", gsu, "exact identifier and representative variant")
    for i, p in enumerate(pairs):
        if p["pair_id"] == pair_id:
            add_map("table2", gsu, "comparison", ROOT / "outcome_data/t06/gwas_gwas_susie_coloc.tsv", f"physical line {i + 2}", f"{pair_id}|G{p['graves_cs']}|A{p['aiht_cs']}", p["classification"])
    selected_locators = {
        "GSU018": {("reeve", "Supplementary Table 1", "31"), ("reeve", "Supplementary Table 3", "8"), ("bujnis", "ST3", "18")},
        "GSU042": {("reeve", "Supplementary Table 1", "65"), ("reeve", "Supplementary Table 3", "240")},
        "GSU052": {("reeve", "Supplementary Table 1", "82"), ("reeve", "Supplementary Table 3", "223"), ("bujnis", "ST4", "188"), ("bujnis", "ST4", "189")},
        "GSU056": {("reeve", "Supplementary Table 1", "84"), ("reeve", "Supplementary Table 3", "225")},
        "GSU073": {("reeve", "Supplementary Table 1", "114"), ("reeve", "Supplementary Table 3", "285")},
        "GSU124": set(),
        "GSU138": {("reeve", "Supplementary Table 1", "204"), ("reeve", "Supplementary Table 3", "396")},
        "GSU174": {("reeve", "Supplementary Table 1", "245"), ("reeve", "Supplementary Table 1", "246"), ("bujnis", "ST3", "108"), ("bujnis", "ST3", "109")} | {("bujnis", "ST4", str(i)) for i in range(1226, 1232)},
        "GSU231": {("reeve", "Supplementary Table 1", "319"), ("reeve", "Supplementary Table 3", "135")},
        "GSU254": {("reeve", "Supplementary Table 1", "345"), ("reeve", "Supplementary Table 3", "159"), ("bujnis", "ST4", "1862"), ("bujnis", "ST4", "1864")},
    }
    for i, ev in enumerate(source_matches):
        if ev["signal_unit_id"] == gsu:
            if (ev["study"], ev["sheet"], ev["excel_row"]) not in selected_locators[gsu]:
                continue
            source = PREV / ("review_reeve_bujnis/sources/reeve_tables.xlsx" if ev["study"] == "reeve" else "review_reeve_bujnis/sources/bujnis_tables.xlsx")
            add_map("table2", gsu, "prior_evidence", source, f"{ev['sheet']}!row{ev['excel_row']}", ev["source_variant"] or ev["source_rsid"], ev["relation"])
    for ev in nature_hits:
        if ev["GSU"] == gsu:
            add_map("table2", gsu, "prior_evidence", PREV / "review_multiome/supp_table_26.xlsx", f"ST26!row{ev['ST26_row']}", ev["trait"] + "|" + ev["variant"], "exact target-trait max-PIP representative")
    matrix_line = next(i + 2 for i, u in enumerate(inventory) if u["signal_unit_id"] == gsu)
    add_map("table2", gsu, "increment", PREV / "locus_increment_matrix.tsv", f"physical line {matrix_line}", gsu, "approved source-bounded interpretation; compressed in English")
    add_map("table2", gsu, "match_level", PREV / "locus_increment_matrix.tsv", f"physical line {matrix_line}", gsu, "approved evidence-match tier; supported by original workbook rows in prior_evidence map")
for gsu, path, locator in (("GSU052", "root_sources/wallace.xml", "Results: Application to Colocalisation Testing of Autoimmune Thyroid Diseases; Figure 5 discussion"), ("GSU231", "root_sources/makkonen.html", "Results: TSHR variants associate with hypo- and hyperthyroidism but not with other phenotypes in FinnGen; Supplemental Fig. 5A/B"), ("GSU124", "review_multiome/supp_table_20.xlsx", "ST20!row101535"), ("GSU124", "review_multiome/tnrc18_same_cell_feature_trait_R12.tsv", "all 4 rows; original source line locators retained")):
    add_map("table2", gsu, "prior_evidence", PREV / path, locator, gsu, "additional original source context")
add_map("table2", "GSU073", "prior_evidence", PREV / "review_reeve_bujnis/sources/bujnis_tables.xlsx", "ST4!row513; column csCount", "LPP; csCount=5", "original published count of regional credible sets")

# Verify consequential literature statements against actual publisher workbooks.
rw = openpyxl.load_workbook(PREV / "review_reeve_bujnis/sources/reeve_tables.xlsx", read_only=True, data_only=True)
rr = list(rw["Supplementary Table 3"].values)
rh = rr[0]
for gsu, rownum in (("GSU018", 8), ("GSU042", 240), ("GSU052", 223), ("GSU056", 225)):
    source = dict(zip(rh, rr[rownum - 1]))
    check("original_Reeve_exact_" + gsu, "chr" + source["v"].replace(":", "_") == units[gsu]["representative_variant"], {"sheet": "Supplementary Table 3", "excel_row": rownum, "variant": source["v"]})
for rownum in (225, 159, 396):
    source = dict(zip(rh, rr[rownum - 1]))
    check("original_Reeve_opposite_" + str(rownum), source["AIHT_beta"] * source["THYTOX_beta"] < 0, {"AIHT_beta": source["AIHT_beta"], "THYTOX_beta": source["THYTOX_beta"]})
for gsu, rownum, traits_expected in (("GSU073", 285, [G, A]), ("GSU138", 396, [A]), ("GSU231", 135, [A]), ("GSU254", 159, [A])):
    source = dict(zip(rh, rr[rownum - 1]))
    variant = "chr" + source["v"].replace(":", "_")
    pair_id = next(r["pair_id"] for r in inventory if r["signal_unit_id"] == gsu)
    for trait in traits_expected:
        regions = {p["graves_region" if trait == G else "aiht_region"] for p in pairs if p["pair_id"] == pair_id}
        matches = [m for m in members if m["trait_id"] == trait and m["region"] in regions and m["cs"] == "1" and m["rsid"] == variant]
        check("original_Reeve_membership_" + gsu + "_" + trait, len(matches) == 1, {"variant": variant, "matched_members": len(matches), "current_cs": 1})
        for i, member in enumerate(members):
            if member in matches:
                add_map("table2", gsu, "match_level", ROOT / "outcome_data/t06/susie_credible_set_members.tsv.gz", f"decompressed physical line {i + 2}", "|".join([trait, member["region"], member["cs"], variant]), "different prior index is a current 95% CS member")
rw.close()

bw = openpyxl.load_workbook(PREV / "review_reeve_bujnis/sources/bujnis_tables.xlsx", read_only=True, data_only=True)
br = list(bw["ST4"].values)
check("original_Bujnis_CTLA4_two_CSs", br[187][3] == 2 and br[188][3] == 2 and br[187][16] == "rs231726" and br[188][16] == "rs12464033", {"ST4_rows": [188, 189], "published_csCount": [br[187][3], br[188][3]]})
check("original_Bujnis_LPP_five_CSs", br[512][3] == 5 and br[512][17] == "LPP", {"ST4_row": 513, "published_csCount": br[512][3]})
bw.close()

nw = openpyxl.load_workbook(PREV / "review_multiome/supp_table_26.xlsx", read_only=True, data_only=True)
nr = list(nw["ST26"].values)
nh = nr[2]
for gsu, rownum, trait in (("GSU052", 5016, A), ("GSU052", 13947, G), ("GSU056", 5020, A), ("GSU124", 5081, A), ("GSU124", 13951, G), ("GSU138", 13952, G), ("GSU174", 4901, A), ("GSU231", 13941, G)):
    source = dict(zip(nh, nr[rownum - 1]))
    check("original_Kanai_exact_" + str(rownum), source["trait"] == trait and source["max_pip_variant"] == units[gsu]["representative_variant"], {"sheet": "ST26", "excel_row": rownum, "trait": source["trait"], "variant": source["max_pip_variant"]})
nw.close()

rp06 = [p for p in pairs if p["pair_id"] == "RP06"]
check("CTLA4_all_five_unresolved_retained", sum(p["classification"].startswith("unresolved") for p in rp06) == 5 and table2_rows[2]["comparison"].count("G2–A") == 4, table2_rows[2]["comparison"])
rt_leads = {r["lead_variant"] for r in catalog if (r["trait_id"] == G and r["region"] == "chr10:60712113-63712113" and r["cs"] == "1") or (r["trait_id"] == A and r["region"] == "chr10:60520112-63520112" and r["cs"] == "2")}
check("RTKN2_lead_representative_distinction", rt_leads == {"chr10_62212113_A_T", "chr10_62286595_T_G"} and units["GSU174"]["representative_variant"] not in rt_leads, sorted(rt_leads))
check("table2_current_all_34_comparisons_mapped", len([r for r in source_maps["table2"] if r["column"] == "comparison"]) == 34, "Includes 8 H4, 11 H3 and 15 unresolved comparisons in the 10 selected region pairs")
check("representative_equivalence_bound_sensitivity", units["GSU138"]["classification_robust_all_bounds"] == "FALSE" and units["GSU231"]["classification_robust_all_bounds"] == "TRUE", {gsu: units[gsu]["classification_across_bounds"] for gsu in ("GSU138", "GSU231")})

payloads = {
    "table1": {"id": "Table 1", "title": "Study datasets, phenotype definitions and comparison units", "columns": table1_columns, "rows": table1_rows, "notes": table1_notes, "references": [r for r in references if r["id"] in {"Kanai 2026", "Sakaue 2021"}]},
    "table2": {"id": "Table 2", "title": "Prior evidence and component-level information added by the present comparison", "columns": table2_columns, "rows": table2_rows, "notes": table2_notes, "references": [r for r in references if r["id"] != "Sakaue 2021"]},
}
payloads["table2"]["source_notes"] = [
    {"row_id": row["signal_unit_id"], "locators": [
        {"source_file": item["source_file"], "source_locator": item["source_locator"], "evidence_key": item["evidence_key"], "relation": item["relation"]}
        for item in source_maps["table2"] if item["row_id"] == row["signal_unit_id"] and item["column"] in {"prior_evidence", "match_level"} and "locus_increment_matrix" not in item["source_file"]
    ]} for row in inventory
]
for name, payload in payloads.items():
    save_json(TABLES / (name + ".json"), payload)
    keys = [c["key"] for c in payload["columns"]]
    write_tsv(TABLES / (name + ".tsv"), ["row_id"] + (["section"] if name == "table1" else []) + keys, payload["rows"])
    map_columns = ["table", "row_id", "column", "source_file", "source_locator", "evidence_key", "relation"]
    write_tsv(MAPS / (name + "_source_map.tsv"), map_columns, source_maps[name])
save_json(TABLES / "references.json", references)

notes = ["# Main tables and table notes", "", "This package presents existing results for author review. No GWAS, fine-mapping, colocalization, enrichment, external validation or equivalence tests were rerun.", ""]
for payload in payloads.values():
    notes += [f"## {payload['id']}. {payload['title']}", ""]
    notes += [f"{i + 1}. {note}" for i, note in enumerate(payload["notes"])] + [""]
notes += ["## References", ""] + [r["citation"] + " https://doi.org/" + r["doi"] for r in references]
notes += ["", "## Source-verification boundaries for the review package", "", "Reeve and Bujnis claims are grounded in the retrieved original publisher workbooks. Kanai 2026 refers to the formal Nature version of record and original supplementary workbooks; the retrieved article HTML is a subscription preview, and full-text access is not claimed. Wallace original XML and Makkonen original publisher HTML support the cited Results statements. The present tables make no first-study or absence-of-prior-literature claim. Cross-study credible-set identity remains unresolved where only a regional or membership match is available.", "", "The workbook is an editable presentation of these frozen outputs and evidence statements. Editing workbook cells does not update the frozen analysis or regenerate these tables; reproducible refresh uses the retained builders.", ""]
(TABLES / "table_notes.md").write_text("\n".join(notes), encoding="utf-8")

for path in INPUT_PATHS:
    check("unchanged:" + relative(path), sha256(path) == input_hashes[relative(path)], input_hashes[relative(path)])
for name, payload in payloads.items():
    reloaded = read_tsv(TABLES / (name + ".tsv"))
    check(name + "_TSV_roundtrip", len(reloaded) == len(payload["rows"]) and {r["row_id"] for r in reloaded} == {r["row_id"] for r in payload["rows"]}, {"rows": len(reloaded)})
    check(name + "_mapping_paths_exist", all((ROOT / r["source_file"]).is_file() for r in source_maps[name]), {"source_map_records": len(source_maps[name])})

validation = {"created_at_utc": datetime.now(timezone.utc).isoformat(), "scope": "Source identity, displayed definitions and mapping checks; no statistical reanalysis", "all_pass": all(c["passed"] for c in checks), "checks": checks, "input_sha256": input_hashes, "permitted_side_effects": ["tables/", "source_data/table*_source_map.tsv", "reviews/table_build*"], "not_performed": ["Statistical model reruns", "New external literature searches", "Methods or Results drafting", "Independent human review"]}
save_json(REVIEWS / "table_build_validation.json", validation)
print(json.dumps({"tables": {k: len(v["rows"]) for k, v in payloads.items()}, "checks": len(checks), "all_pass": validation["all_pass"], "failures": [c for c in checks if not c["passed"]]}, ensure_ascii=False, indent=2))
if not validation["all_pass"]:
    raise SystemExit(1)
